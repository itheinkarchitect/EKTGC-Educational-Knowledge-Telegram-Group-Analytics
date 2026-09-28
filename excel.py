import os
from openpyxl import Workbook, load_workbook

FILE_NAME = "analytics.xlsx"

HEADERS = [
    "ID сообщения",
    "Дата сообщения",
    "ID пользователя",
    "Имя",
    "Фамилия",
    "Время прочтения"
]


def open_excel():
    if os.path.exists(FILE_NAME):
        workbook = load_workbook(FILE_NAME)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Прочтения"

        sheet.append(HEADERS)

        workbook.save(FILE_NAME)

    return workbook


def save_readings(message, readers):
    workbook = open_excel()
    sheet = workbook["Прочтения"]

    existing_records = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        message_id = row[0]
        user_id = row[2]

        existing_records.add((message_id, user_id))

    for reader in readers:
        record = (message.id, reader["user_id"])

        if record in existing_records:
            continue

        sheet.append([
            message.id,
            message.date,
            reader["user_id"],
            reader["first_name"],
            reader["last_name"],
            reader["read_at"]
        ])

        existing_records.add(record)

    workbook.save(FILE_NAME)